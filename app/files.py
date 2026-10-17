"""Value-table import/export. Never overwrite the imported workbook implicitly."""
from __future__ import annotations

import csv
import io
import os
import re
import tempfile
import zipfile
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .data import Column, Record, Sheet, infer_settings

MAX_ROWS = 50_000
MAX_COLUMNS = 200
MAX_FILE_BYTES = 50 * 1024 * 1024


def text_value(value, number_format="") -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (float, int)) and re.fullmatch(r"0{2,20}", number_format or ""):
        if float(value).is_integer():
            return str(int(value)).zfill(len(number_format))
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def unique_headers(values) -> list[str]:
    used, headers = set(), []
    for index, value in enumerate(values):
        base = str(value or "").strip() or f"列{index + 1}"
        label, suffix = base, 2
        while label in used:
            label = f"{base} ({suffix})"
            suffix += 1
        used.add(label)
        headers.append(label)
    return headers


def from_matrix(title: str, matrix) -> Sheet | None:
    iterator = iter(matrix)
    header = next(iterator, None)
    if header is None:
        return None
    if len(header) > MAX_COLUMNS:
        raise ValueError(f"最多支持 {MAX_COLUMNS} 列，请先拆分表格")
    columns = [Column(label, "multiline" if label in {"反馈", "反馈记录", "催办回复", "说明", "免催原因"} else "text")
               for label in unique_headers(header)]
    if not columns:
        return None
    sheet = Sheet(title, columns, settings=infer_settings(columns))
    for values in iterator:
        if len(values) > MAX_COLUMNS:
            raise ValueError(f"最多支持 {MAX_COLUMNS} 列")
        if len(values) > len(columns) and any(str(v or "") for v in values[len(columns):]):
            raise ValueError("数据行比表头更宽，请为每个有数据的列提供表头")
        if len(sheet.rows) >= MAX_ROWS:
            raise ValueError(f"每张表最多支持 {MAX_ROWS} 条记录，请先拆分表格")
        if any("\0" in str(value) or len(str(value)) > 32_000 for value in values):
            raise ValueError("单元格包含空字符或超过 32,000 字，无法安全导入")
        # Retain blank rows too: source ordering and row boundaries are meaningful.
        sheet.rows.append(Record({col.id: str(values[i] if i < len(values) else "")
                                  for i, col in enumerate(columns)}))
    return sheet


def import_file(path: Path) -> list[Sheet]:
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("文件超过 50 MB，请先拆分后导入")
    suffix = path.suffix.lower()
    sheets = []
    if suffix == ".xlsx":
        with zipfile.ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024:
                raise ValueError("工作簿解压后过大，请先拆分后导入")
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            for worksheet in workbook.worksheets:
                if worksheet.max_column > MAX_COLUMNS or worksheet.max_row > MAX_ROWS + 1:
                    raise ValueError(f"工作表「{worksheet.title}」超过 50,000 行或 200 列限制")
                matrix = ([text_value(cell.value, cell.number_format) for cell in row]
                          for row in worksheet.iter_rows())
                sheet = from_matrix(worksheet.title, matrix)
                if sheet:
                    sheets.append(sheet)
        finally:
            workbook.close()
    elif suffix == ".csv":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("CSV 编码无法识别，请另存为 UTF-8 CSV")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        sheet = from_matrix(path.stem, csv.reader(io.StringIO(text, newline=""), dialect))
        if sheet:
            sheets.append(sheet)
    else:
        raise ValueError("支持 .xlsx 和 .csv；旧版 .xls 请先另存为 .xlsx")
    if not sheets:
        raise ValueError("没有可导入的表头或工作表")
    return sheets


def _excel_text(worksheet, row, column, value):
    # User content that starts with '=' remains literal text, never a new formula.
    cell = worksheet.cell(row, column)
    cell.value = value
    cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return cell


def _csv_text(value: str) -> str:
    # CSV has no cell types. Escape formula-triggering content for Office clients.
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def export_file(path: Path, sheets: list[Sheet]):
    if not sheets:
        raise ValueError("没有可导出的表格")
    suffix = path.suffix.lower()
    if suffix not in (".xlsx", ".csv"):
        raise ValueError("请选择 .xlsx 或 .csv 文件")
    if suffix == ".csv" and len(sheets) != 1:
        raise ValueError("CSV 仅支持单表，请选择当前表导出或使用 XLSX")
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temp_name = tempfile.mkstemp(prefix=".table-export-", suffix=suffix, dir=path.parent)
    os.close(descriptor)
    temporary = Path(temp_name)
    try:
        if suffix == ".csv":
            sheet = sheets[0]
            with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow([_csv_text(col.label) for col in sheet.columns])
                for row in sheet.rows:
                    writer.writerow([_csv_text(row.cells.get(col.id, "")) for col in sheet.columns])
        else:
            workbook = Workbook()
            workbook.remove(workbook.active)
            for sheet in sheets:
                title = re.sub(r"[\\/*?:\[\]]", "_", sheet.title).strip("'")[:31] or "表格"
                worksheet = workbook.create_sheet(title)
                worksheet.freeze_panes = "A2"
                worksheet.auto_filter.ref = f"A1:{get_column_letter(len(sheet.columns))}{len(sheet.rows) + 1}"
                for column, definition in enumerate(sheet.columns, 1):
                    cell = _excel_text(worksheet, 1, column, definition.label)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="2563EB")
                    worksheet.column_dimensions[get_column_letter(column)].width = 22
                for number, row in enumerate(sheet.rows, 2):
                    for column, definition in enumerate(sheet.columns, 1):
                        _excel_text(worksheet, number, column, row.cells.get(definition.id, ""))
            try:
                workbook.save(temporary)
            finally:
                workbook.close()
        # Atomic replacement: a failed export leaves the existing destination intact.
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
