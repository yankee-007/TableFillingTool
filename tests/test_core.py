from __future__ import annotations

import csv
import sqlite3
import ctypes
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app.data import Column, Record, Sheet, Store, infer_settings
from app.files import export_file, from_matrix, import_file, unique_headers
from app.matching import candidates, name_in_title
from app.service import Workspace
from app.wecom import ContactSearch, NativeDesktop, Window


def fixture(title="一班", names=("张三", "李四"), prefix="py169"):
    columns = [Column("姓名"), Column("联系人备注"), Column("反馈记录", "multiline")]
    sheet = Sheet(title, columns, settings=infer_settings(columns))
    sheet.settings["prefix"] = prefix
    sheet.rows = [Record({columns[0].id: name, columns[1].id: "", columns[2].id: ""}) for name in names]
    return sheet


class FileTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_headers_are_unique_even_when_suffixes_already_exist(self):
        self.assertEqual(unique_headers(["姓名", "姓名", "姓名 (2)", ""]),
                         ["姓名", "姓名 (2)", "姓名 (2) (2)", "列4"])

    def test_xlsx_all_sheets_leading_zero_formula_and_source_unchanged(self):
        path = self.root / "source.xlsx"
        book = Workbook()
        ws = book.active
        ws.title = "一班"
        ws.append(["学号", "姓名", "反馈"])
        ws.append([12, "张三", "=1+2"])
        ws["A2"].number_format = "00000"
        ws.append(["0013", "李四", "第一行\n第二行"])
        book.create_sheet("二班").append(["姓名", "备注"])
        book.save(path)
        before = path.read_bytes()
        sheets = import_file(path)
        self.assertEqual(len(sheets), 2)
        self.assertEqual(list(sheets[0].rows[0].cells.values()), ["00012", "张三", "=1+2"])
        target = self.root / "output.xlsx"
        export_file(target, sheets)
        result = load_workbook(target)
        try:
            self.assertEqual(result["一班"]["A2"].value, "00012")
            self.assertEqual(result["一班"]["C2"].value, "=1+2")
            self.assertEqual(result["一班"]["C2"].data_type, "s")
            self.assertEqual(result["一班"]["C3"].value, "第一行\n第二行")
            self.assertEqual(path.read_bytes(), before)
        finally:
            result.close()

    def test_csv_chinese_multiline_comma_and_gb18030(self):
        path = self.root / "名单.csv"
        path.write_bytes('姓名,反馈\r\n张三,"中文,逗号\n下一行"\r\n'.encode("gb18030"))
        sheet = import_file(path)[0]
        self.assertEqual(list(sheet.rows[0].cells.values()), ["张三", "中文,逗号\n下一行"])
        target = self.root / "out.csv"
        export_file(target, [sheet])
        reloaded = import_file(target)[0]
        self.assertEqual(list(reloaded.rows[0].cells.values()), list(sheet.rows[0].cells.values()))

    def test_csv_formula_safety_and_no_multi_sheet_silent_loss(self):
        sheet = from_matrix("Sheet", [["姓名", "反馈"], ["张三", "=HYPERLINK(1)"]])
        path = self.root / "out.csv"
        export_file(path, [sheet])
        with path.open(encoding="utf-8-sig", newline="") as handle:
            self.assertEqual(list(csv.reader(handle))[1][1], "'=HYPERLINK(1)")
        with self.assertRaises(ValueError):
            export_file(path, [sheet, fixture()])

    def test_export_failure_keeps_existing_file_and_cleans_temporary(self):
        path = self.root / "old.xlsx"
        path.write_bytes(b"existing content")
        with patch("app.files.os.replace", side_effect=PermissionError("locked")):
            with self.assertRaises(PermissionError):
                export_file(path, [fixture()])
        self.assertEqual(path.read_bytes(), b"existing content")
        self.assertEqual(list(self.root.glob(".table-export-*")), [])

    def test_unsupported_or_missing_header_is_explicit(self):
        path = self.root / "legacy.xls"
        path.write_bytes(b"dummy")
        with self.assertRaisesRegex(ValueError, "xlsx"):
            import_file(path)
        with self.assertRaisesRegex(ValueError, "表头"):
            from_matrix("bad", [["姓名"], ["张三", "hidden value"]])

    def test_oversize_cell_or_embedded_null_is_rejected(self):
        for value in ("a" * 32_001, "bad\0cell"):
            with self.assertRaises(ValueError):
                from_matrix("invalid", [["姓名"], [value]])


class PersistenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "tables.sqlite3"
        self.store = Store(self.path)
        self.workspace = Workspace(self.store)
        self.sheet = fixture()
        self.assertTrue(self.workspace.add_sheets([self.sheet]))

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def test_restart_retains_ids_values_settings_and_new_columns(self):
        row = self.sheet.rows[0]
        cid = self.sheet.columns[2].id
        self.assertTrue(self.workspace.save_values(self.sheet.id, row.id, {cid: "已联系；下次沟通"}))
        new_column = Column("日期", "date")
        self.assertTrue(self.workspace.add_column(self.sheet.id, new_column))
        self.assertTrue(self.workspace.save_values(self.sheet.id, row.id, {new_column.id: "2026-10-17"}))
        settings = dict(self.sheet.settings, prefix="新前缀")
        self.assertTrue(self.workspace.configure(self.sheet.id, settings))
        restored = self.store.load()[0]
        self.assertEqual(restored.id, self.sheet.id)
        self.assertEqual(restored.rows[0].id, row.id)
        self.assertEqual(restored.rows[0].cells[cid], "已联系；下次沟通")
        self.assertEqual(restored.rows[0].cells[new_column.id], "2026-10-17")
        self.assertEqual(restored.settings["prefix"], "新前缀")

    def test_invalid_identity_and_invalid_date_never_mutate_another_record(self):
        cid = self.sheet.columns[2].id
        self.assertFalse(self.workspace.save_values(self.sheet.id, "stale-record", {cid: "错误内容"}))
        self.assertEqual(self.sheet.rows[1].cells[cid], "")
        column = Column("日期", "date")
        self.workspace.add_column(self.sheet.id, column)
        self.assertFalse(self.workspace.save_values(self.sheet.id, self.sheet.rows[0].id,
                                                    {cid: "must be atomic", column.id: "not-a-date"}))
        self.assertEqual(self.sheet.rows[0].cells[cid], "")

    def test_failed_save_does_not_claim_saved_or_mutate_memory(self):
        cid = self.sheet.columns[2].id
        with patch.object(self.store, "save_cells", side_effect=OSError("disk full")):
            self.assertFalse(self.workspace.save_values(self.sheet.id, self.sheet.rows[0].id, {cid: "未落盘"}))
        self.assertEqual(self.sheet.rows[0].cells[cid], "")
        self.assertEqual(self.store.load()[0].rows[0].cells[cid], "")

    def test_import_transaction_rolls_back_all_sheets(self):
        new_sheet = fixture("new")
        collision = fixture("collision")
        collision.id = self.sheet.id
        self.assertFalse(self.workspace.add_sheets([new_sheet, collision]))
        self.assertEqual(len(self.store.load()), 1)
        self.assertEqual(len(self.workspace.sheets), 1)

    def test_duplicate_import_uses_new_identity_and_unique_table_name(self):
        another = fixture()
        self.assertTrue(self.workspace.add_sheets([another]))
        self.assertEqual(another.title, "一班 (2)")
        self.assertNotEqual(another.id, self.sheet.id)
        self.assertEqual(len(self.store.load()), 2)


class MatchingTests(unittest.TestCase):
    def test_han_boundaries_and_main_window_excluded(self):
        self.assertFalse(name_in_title("张三", "张三丰"))
        self.assertFalse(name_in_title("张三", "老张三"))
        self.assertTrue(name_in_title("张三", "py169张三 - 企业微信"))
        self.assertFalse(name_in_title("WeCom", "WeCom"))
        self.assertFalse(name_in_title("Ann", "Anna"))
        self.assertTrue(name_in_title("Ann", "Ann - WeCom"))

    def test_prefix_priority_duplicate_names_and_table_scope(self):
        one, two = fixture("一班", ("张三",), "py169"), fixture("二班", ("张三",), "py170")
        self.assertEqual(candidates([one, two], "py170张三"), [(two.id, two.rows[0].id)])
        self.assertEqual(len(candidates([one, two], "张三")), 2)
        self.assertEqual(candidates([one, two], "张三", one.id), [(one.id, one.rows[0].id)])
        self.assertEqual(candidates([one, two], "未知姓名"), [])

    def test_explicit_remark_can_identify_contact_without_legal_name(self):
        sheet = fixture(names=("张三",))
        sheet.rows[0].cells[sheet.columns[1].id] = "P2026-小王"
        self.assertEqual(candidates([sheet], "P2026-小王"), [(sheet.id, sheet.rows[0].id)])


class FakeDesktop:
    def __init__(self, chat_title="py169张三", executable=r"C:\WeCom\WXWork.exe"):
        self.main = Window(1, 42, "企业微信", executable)
        self.chat = Window(2, 42, chat_title, executable)
        self.current = self.main
        self.events = []
        self.after_search = None

    def main_window(self):
        return self.main

    def activate(self, window):
        self.current = window

    def foreground(self):
        return self.current

    def clipboard_text(self):
        return "py169张三"

    def keys(self, *keys):
        self.events.append(keys)
        if keys == ("enter",) and self.after_search:
            self.current = self.after_search
        if keys == ("ctrl", "o"):
            self.current = self.chat
        if keys == ("ctrl", "w"):
            self.current = self.main


class ContactTests(unittest.TestCase):
    def search(self, desktop):
        return ContactSearch(desktop, wait=0, sleep=lambda _: None)

    def test_open_contact_never_sends_message(self):
        desktop = FakeDesktop()
        self.assertIn("已打开", self.search(desktop).open("py169张三"))
        self.assertEqual(desktop.events, [("ctrl", "f"), ("ctrl", "a"), ("ctrl", "v"), ("enter",), ("ctrl", "o")])
        self.assertEqual(desktop.current.handle, 2)

    def test_optional_close_verified_chat(self):
        desktop = FakeDesktop()
        self.search(desktop).open("py169张三", keep_float=False)
        self.assertEqual(desktop.events[-1], ("ctrl", "w"))
        self.assertEqual(desktop.current.handle, 1)

    def test_unverified_search_does_not_open_float(self):
        desktop = FakeDesktop()
        self.assertIn("未验证", self.search(desktop).open("py169张三", verify=False))
        self.assertNotIn(("ctrl", "o"), desktop.events)

    def test_mismatched_contact_blocks_further_actions(self):
        desktop = FakeDesktop("py169李四")
        with self.assertRaisesRegex(RuntimeError, "不匹配"):
            self.search(desktop).open("py169张三")
        self.assertNotIn(("ctrl", "w"), desktop.events)

    def test_focus_change_and_non_wecom_process_abort(self):
        desktop = FakeDesktop()
        desktop.after_search = Window(99, 88, "unrelated", "unrelated.exe")
        with self.assertRaisesRegex(RuntimeError, "焦点"):
            self.search(desktop).open("py169张三")
        self.assertNotIn(("ctrl", "o"), desktop.events)
        desktop = FakeDesktop(executable="not-wecom.exe")
        with self.assertRaises(RuntimeError):
            self.search(desktop).open("py169张三")
        self.assertEqual(desktop.events, [])

    def test_empty_multiline_and_cancelled_inputs_press_no_keys(self):
        for value in ("", "张三\n李四"):
            desktop = FakeDesktop()
            with self.assertRaises(ValueError):
                self.search(desktop).open(value)
            self.assertEqual(desktop.events, [])
        desktop = FakeDesktop()
        with self.assertRaisesRegex(RuntimeError, "取消"):
            ContactSearch(desktop, wait=0, sleep=lambda _: None, cancelled=lambda: True).open("张三")
        self.assertEqual(desktop.events, [])

    def test_clipboard_change_never_pastes_unrelated_content(self):
        desktop = FakeDesktop()
        desktop.clipboard_text = lambda: "unrelated clipboard"
        with self.assertRaisesRegex(RuntimeError, "剪贴板"):
            self.search(desktop).open("py169张三")
        self.assertNotIn(("ctrl", "v"), desktop.events)
        self.assertNotIn(("enter",), desktop.events)

    @unittest.skipUnless(sys.platform == "win32", "Windows ABI only; no actual keyboard input")
    def test_native_input_abi_and_modifier_release_with_sendinput_mocked(self):
        desktop = NativeDesktop()
        events = []

        def fake_send(count, array, size):
            events.append((size, [(array[i].keyboard.vk, array[i].keyboard.flags) for i in range(count)]))
            return count

        with patch.object(desktop.user, "SendInput", fake_send):
            desktop.keys("ctrl", "v")
        self.assertEqual(events[0][0], 40 if ctypes.sizeof(ctypes.c_void_p) == 8 else 28)
        self.assertEqual(events[0][1], [(0x11, 0), (0x56, 0), (0x56, 2), (0x11, 2)])


if __name__ == "__main__":
    unittest.main()
